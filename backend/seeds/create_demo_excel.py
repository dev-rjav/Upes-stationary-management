"""
Generates demo_vendor_invoice.xlsx — a sample vendor delivery note for testing
the parse-excel / stock-in upload feature.

Run from the backend/ directory:
    python seeds/create_demo_excel.py

Produces:  seeds/demo_vendor_invoice.xlsx

The file intentionally contains a mix of match types:
  - High confidence (green): names close to canonical
  - Review needed (yellow): moderately fuzzy names
  - Unrecognised (red): items that won't match the catalog
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), 'demo_vendor_invoice.xlsx')

ROWS = [
    # (item_description, hsn_code, qty, unit, rate)
    # --- High confidence matches ---
    ('Ball Point Pen',               '9608',  200, 'EA',   12.00),
    ('A4 Paper Ream 75 GSM',         '4802',   30, 'REAM', 285.00),
    ('Whiteboard Marker',            '9608',   60, 'EA',   54.00),
    ('L Folder',                     '4820',   50, 'EA',   27.00),
    ('Correction Fluid Whitener',    '3824',   40, 'EA',   26.00),
    ('Pencil Natraj',                '9609',  100, 'EA',    8.00),
    ('Sticky Pad 3x3',               '4820',   30, 'PAC',  52.00),
    ('Conference Notepad',           '4820',   25, 'EA',   38.00),
    ('Rubber Bands',                 '4016',   15, 'PKT',  29.00),
    ('Stapler No 10',                '8472',   10, 'EA',   92.00),
    ('Drawing Pin',                  '7317',   20, 'EA',   17.00),
    ('Scissors',                     '8213',   12, 'EA',   53.00),
    # --- Review needed (moderate match) ---
    ('Cello Tape Roll',              '3919',   40, 'EA',   33.00),
    ('Binder 32 MM Clips',           '8305',   25, 'DZ',   29.00),
    ('Stapler Pin Small',            '8305',   50, 'EA',   19.00),
    ('Whiteboard Eraser',            '9603',    8, 'EA',   68.00),
    ('Gum Glue Bottle',              '3506',   15, 'EA',   52.00),
    ('Printer Paper A3',             '4802',   10, 'REAM', 415.00),
    ('Page Flag Markers',            '4820',   20, 'PAC',  43.00),
    # --- Unrecognised items ---
    ('HP LaserJet Toner 85A',        '8443',    5, 'EA', 1850.00),
    ('USB Pendrive 32GB',            '8471',   10, 'EA',  550.00),
    ('Desk Organiser Wooden',        '4421',    6, 'EA',  340.00),
]

HEADER_FILL = PatternFill('solid', fgColor='1F2F55')
HEADER_FONT = Font(color='FFFFFF', bold=True, size=11)
SUBHEADER_FILL = PatternFill('solid', fgColor='EEF0F5')
TOTAL_FILL = PatternFill('solid', fgColor='F0F9F4')
THIN = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='thin', color='DDDDDD'),
)

COL_WIDTHS = [6, 42, 12, 10, 10, 16, 16]
COL_HEADERS = ['S.No', 'Item Description', 'HSN Code', 'Quantity', 'Unit', 'Rate/Unit (₹)', 'Amount (₹)']


def make_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Vendor Invoice'

    # Title block
    ws.merge_cells('A1:G1')
    ws['A1'] = 'QUALITY STATIONERY SUPPLIES'
    ws['A1'].font = Font(bold=True, size=14, color='1F2F55')
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:G2')
    ws['A2'] = 'Delivery Challan / Invoice'
    ws['A2'].font = Font(italic=True, size=11, color='555555')
    ws['A2'].alignment = Alignment(horizontal='center')

    ws['A4'] = 'Invoice No:'
    ws['B4'] = 'QSS-2026-1007'
    ws['B4'].font = Font(bold=True)
    ws['D4'] = 'Date:'
    ws['E4'] = date.today().strftime('%d-%b-%Y')
    ws['E4'].font = Font(bold=True)

    ws['A5'] = 'Supplied to:'
    ws['B5'] = 'UPES — Stationery Department'
    ws['B5'].font = Font(bold=True)

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[6].height = 4  # spacer

    # Column headers at row 7
    HEADER_ROW = 7
    for col_idx, (header, width) in enumerate(zip(COL_HEADERS, COL_WIDTHS), start=1):
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = THIN
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[HEADER_ROW].height = 20

    # Data rows
    for i, (desc, hsn, qty, unit, rate) in enumerate(ROWS, start=1):
        row = HEADER_ROW + i
        amount = qty * rate
        row_fill = PatternFill('solid', fgColor='FAFAFA') if i % 2 == 0 else None

        values = [i, desc, hsn, qty, unit, rate, amount]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.border = THIN
            cell.alignment = Alignment(
                horizontal='center' if col_idx != 2 else 'left',
                vertical='center'
            )
            if row_fill:
                cell.fill = row_fill
            if col_idx in (6, 7):
                cell.number_format = '#,##0.00'

    # Totals row
    total_row = HEADER_ROW + len(ROWS) + 1
    ws.cell(row=total_row, column=1, value='').border = THIN
    ws.merge_cells(f'B{total_row}:F{total_row}')
    total_label = ws.cell(row=total_row, column=2, value='TOTAL')
    total_label.font = Font(bold=True)
    total_label.alignment = Alignment(horizontal='right', vertical='center')
    total_label.border = THIN
    total_label.fill = TOTAL_FILL

    total_amt = sum(r[2] * r[4] for r in ROWS)
    total_cell = ws.cell(row=total_row, column=7, value=total_amt)
    total_cell.font = Font(bold=True)
    total_cell.number_format = '#,##0.00'
    total_cell.fill = TOTAL_FILL
    total_cell.border = THIN
    total_cell.alignment = Alignment(horizontal='center', vertical='center')

    # Footer note
    note_row = total_row + 2
    ws.merge_cells(f'A{note_row}:G{note_row}')
    note = ws.cell(
        row=note_row, column=1,
        value='Note: Upload this file via Stock-In → Upload Excel to test the fuzzy-matching parser.'
    )
    note.font = Font(italic=True, size=9, color='888888')
    note.alignment = Alignment(horizontal='left')

    wb.save(OUTPUT_PATH)
    print(f'Created: {OUTPUT_PATH}')
    print(f'  {len(ROWS)} rows — ~12 high-confidence, ~7 review-needed, ~3 unrecognised')


if __name__ == '__main__':
    make_excel()
