"""Dashboard endpoint + report generation (CSV/XLSX) and fulfilled-only accounting."""
import io
from openpyxl import load_workbook
from app.services.report_service import generate_csv, generate_xlsx

AREA = 'Dashboard & Reports'


def _fulfill(client, sap, item_id, qty):
    r = client.post('/api/requests', json={
        'sap_id': sap, 'cluster': 'ENGINEERING', 'items': [{'item_id': item_id, 'qty': qty}],
    })
    client.patch(f"/api/requests/{r.get_json()['id']}/approve")
    return r


def test_dashboard_requires_login(seed, client):
    """Dashboard is receptionist-only."""
    assert client.get('/api/dashboard').status_code == 401


def test_dashboard_returns_expected_keys(seed, as_admin):
    """Dashboard returns the KPI/chart keys the frontend expects."""
    _fulfill(as_admin, '40000001', seed['pen_id'], 3)
    body = as_admin.get('/api/dashboard').get_json()
    for key in ('total_spend', 'total_units', 'monthly_trend', 'top_qty', 'low_stock', 'teacher_usage', 'activity'):
        assert key in body


def test_department_breakdown_requires_cluster(seed, as_admin):
    """The department breakdown needs a cluster parameter."""
    assert as_admin.get('/api/dashboard/department').status_code == 400


def test_csv_report_has_header(seed, as_admin):
    """CSV export includes the expected header row."""
    _fulfill(as_admin, '40000001', seed['pen_id'], 2)
    data = generate_csv({}).decode('utf-8')
    assert 'Date' in data and 'SAP ID' in data and 'Total' in data


def test_report_counts_only_fulfilled(seed, as_admin):
    """Pending requests are excluded from reports; only fulfilled lines count."""
    # one fulfilled, one left pending
    _fulfill(as_admin, '40000001', seed['pen_id'], 2)
    as_admin.post('/api/requests', json={
        'sap_id': '40000002', 'cluster': 'ENGINEERING', 'items': [{'item_id': seed['pen_id'], 'qty': 4}],
    })
    data = generate_csv({}).decode('utf-8')
    # 1 header + exactly 1 data line
    non_empty = [ln for ln in data.splitlines() if ln.strip()]
    assert len(non_empty) == 2


def test_xlsx_has_five_sheets(seed, as_admin):
    """The Excel report contains the five expected sheets."""
    _fulfill(as_admin, '40000001', seed['pen_id'], 2)
    wb = load_workbook(io.BytesIO(generate_xlsx({})))
    assert wb.sheetnames == ['Stationery Issued', 'Departmental Expenditure',
                             'Departmental Detail', 'Department Summary', 'MIS']


def test_xlsx_house_style_applied(seed, as_admin):
    """Headers use the yellow fill and prices use dash-for-zero accounting format."""
    _fulfill(as_admin, '40000001', seed['pen_id'], 2)
    wb = load_workbook(io.BytesIO(generate_xlsx({})))
    ws = wb['Departmental Expenditure']
    # header band yellow
    assert ws['A2'].fill.fgColor.rgb.endswith('FFFF00')
    # a price cell uses an accounting-style number format (contains the dash token)
    assert '"-"' in ws['D4'].number_format
