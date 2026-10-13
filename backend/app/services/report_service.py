import csv
import io
from datetime import datetime

from sqlalchemy import text, func

MONTHS_SHORT = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']


def _apply_window(query, filters):
    """Apply the start/end fulfilled-date window to a query on Request."""
    from ..models import Request
    if filters.get('start'):
        try:
            query = query.filter(Request.fulfilled_at >= datetime.fromisoformat(filters['start']))
        except ValueError:
            pass
    if filters.get('end'):
        try:
            query = query.filter(Request.fulfilled_at <= datetime.fromisoformat(filters['end']))
        except ValueError:
            pass
    return query


def _rate_map():
    """{item_id: latest stock-in rate}. One query instead of per-row lookups."""
    from ..models import StockIn
    rates = {}
    for si in StockIn.query.order_by(StockIn.id.asc()).all():
        rates[si.item_id] = si.rate_per_unit  # ascending => last write wins
    return rates


def _get_report_rows(filters):
    from ..models import Request, RequestLine, StationaryItem, Teacher, Cluster, Receptionist
    from .. import db

    query = (
        db.session.query(Request, RequestLine, StationaryItem)
        .join(RequestLine, RequestLine.request_id == Request.id)
        .join(StationaryItem, RequestLine.item_id == StationaryItem.id)
        .filter(Request.status == 'fulfilled')
    )

    if filters.get('start'):
        try:
            query = query.filter(Request.fulfilled_at >= datetime.fromisoformat(filters['start']))
        except ValueError:
            pass

    if filters.get('end'):
        try:
            query = query.filter(Request.fulfilled_at <= datetime.fromisoformat(filters['end']))
        except ValueError:
            pass

    # Join Teacher once for both department and sap_id filters
    if filters.get('cluster') or filters.get('sap_id'):
        query = query.join(Teacher, Request.teacher_id == Teacher.id)
        if filters.get('sap_id'):
            query = query.filter(Teacher.sap_id == filters['sap_id'])
        if filters.get('cluster'):
            query = query.join(Cluster, Teacher.cluster_id == Cluster.id)
            query = query.filter(Cluster.name == filters['cluster'])

    if filters.get('item_id'):
        query = query.filter(RequestLine.item_id == int(filters['item_id']))

    if filters.get('actioned_by') and filters['actioned_by'] != 'All':
        query = (query
                 .join(Receptionist, Request.fulfilled_by == Receptionist.id)
                 .filter(Receptionist.display_name == filters['actioned_by']))

    return query.all()


def generate_csv(filters):
    rows = _get_report_rows(filters)
    rates = _rate_map()
    output = io.StringIO()
    writer = csv.writer(output)
    # Rupee sign appears on the price column HEADERS only, never against each value.
    writer.writerow(['Date', 'Ref', 'User', 'SAP ID', 'Department', 'Item', 'Qty', 'Rate (₹)', 'Total (₹)', 'Actioned By'])

    for req, line, item in rows:
        rate = rates.get(item.id, 0.0)
        qty = line.quantity_fulfilled or 0
        total = qty * rate
        fulfilled_date = req.fulfilled_at.strftime('%Y-%m-%d') if req.fulfilled_at else ''
        actioned = req.fulfiller.display_name if req.fulfiller else ''
        dept = req.teacher.cluster.name if req.teacher and req.teacher.cluster else ''
        sap_id = req.teacher.sap_id if req.teacher else ''
        user_name = req.teacher.name if req.teacher else ''
        writer.writerow([
            fulfilled_date, req.ref_number, user_name, sap_id, dept,
            item.name, qty, f'{rate:.2f}', f'{total:.2f}', actioned,
        ])

    return output.getvalue().encode('utf-8')


def generate_pdf(filters):
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
    except ImportError:
        raise RuntimeError('reportlab not installed')

    rows = _get_report_rows(filters)
    rates = _rate_map()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=20, rightMargin=20, topMargin=30, bottomMargin=20)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph('Stationery Issuance Report', styles['Heading1']))
    elements.append(Spacer(1, 12))

    # Rupee sign is in the header cells only; values are plain numbers.
    table_data = [['Date', 'Ref', 'User', 'SAP ID', 'Department', 'Item', 'Qty', 'Rate (₹)', 'Total (₹)', 'By']]
    for req, line, item in rows:
        rate = rates.get(item.id, 0.0)
        qty = line.quantity_fulfilled or 0
        total = qty * rate
        fulfilled_date = req.fulfilled_at.strftime('%d/%m/%Y') if req.fulfilled_at else ''
        actioned = req.fulfiller.display_name if req.fulfiller else ''
        dept = req.teacher.cluster.name if req.teacher and req.teacher.cluster else ''
        table_data.append([
            fulfilled_date, req.ref_number, req.teacher.name if req.teacher else '',
            req.teacher.sap_id if req.teacher else '', dept,
            item.name, str(qty),
            f'{rate:.2f}', f'{total:.2f}', actioned,
        ])

    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F2F55')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F7F7F5')]),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#E5E5E5')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)
    doc.build(elements)
    return buffer.getvalue()


# ----------------------------------------------------------------------------
# Polished Excel export — mirrors the manual reference workbooks in /sheets,
# rebuilt live from fulfilled requests. House style (from the reference files):
#   * yellow (FFFF00) bold header band, black text, thin borders
#   * cream-filled "Item Price" column
#   * accounting number formats so zero renders as a dash "-"
#   * bold TOTAL rows
# Sheets:
#   1. "Stationery Issued"        -> per-issuance log
#   2. "Departmental Expenditure" -> items x departments qty/total matrix
#   3. "Departmental Detail"      -> one stacked block per department
#   4. "Department Summary"       -> spend per department
#   5. "MIS"                      -> items x months spend overview
# ----------------------------------------------------------------------------

# Accounting formats — a 0 shows as "-", negatives in ( ), matching the originals.
_FMT_MONEY = r'_(* #,##0.00_);_(* \(#,##0.00\);_(* "-"??_);_(@_)'
_FMT_INT = r'_(* #,##0_);_(* \(#,##0\);_(* "-"??_);_(@_)'

_CLR_YELLOW = 'FFFF00'   # header band
_CLR_CREAM = 'FFF3CF'    # item-price / fixed columns
_CLR_TOTAL = 'FFF9E3'    # total row tint


def _xl_styles():
    """Build the reusable openpyxl style objects (one place, house style)."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    thin = Side(style='thin', color='BFBFBF')
    return {
        'border': Border(left=thin, right=thin, top=thin, bottom=thin),
        'header_font': Font(bold=True, color='000000', size=11),
        'title_font': Font(bold=True, size=14),
        'bold': Font(bold=True),
        'yellow': PatternFill('solid', fgColor=_CLR_YELLOW),
        'cream': PatternFill('solid', fgColor=_CLR_CREAM),
        'total': PatternFill('solid', fgColor=_CLR_TOTAL),
        'center': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'left': Alignment(horizontal='left', vertical='center'),
    }

def _report_year(filters):
    """Pick a year for the MIS/monthly sheet from the filter window."""
    for key in ('start', 'end'):
        val = filters.get(key)
        if val:
            try:
                return datetime.fromisoformat(val).year
            except ValueError:
                pass
    return datetime.now().year


def _period_label(filters):
    """Human label for the filtered window, e.g. 'September 2026' or '2026'."""
    start = filters.get('start')
    if start:
        try:
            d = datetime.fromisoformat(start)
            return d.strftime('%B %Y')
        except ValueError:
            pass
    return str(_report_year(filters))


def generate_xlsx(filters):
    """Polished multi-sheet workbook rebuilt live from fulfilled requests.

    Styling mirrors the manual reference workbooks in /sheets (yellow header
    band, cream price column, dash-for-zero accounting formats, bold TOTAL rows).
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from ..models import Request, RequestLine, StationaryItem, Cluster, Teacher
    from .. import db

    S = _xl_styles()
    rates = _rate_map()
    rows = _get_report_rows(filters)
    period = _period_label(filters)

    def head(cell, fill='yellow'):
        cell.font = S['header_font']
        cell.fill = S[fill]
        cell.alignment = S['center']
        cell.border = S['border']

    def box(cell, numfmt=None, bold=False, fill=None, align='left'):
        cell.border = S['border']
        cell.alignment = S['left'] if align == 'left' else S['center']
        if numfmt:
            cell.number_format = numfmt
        if bold:
            cell.font = S['bold']
        if fill:
            cell.fill = S[fill]
        return cell

    wb = Workbook()

    # Departments in canonical seed order; items by HSN then name (catalog order).
    departments = Cluster.query.order_by(Cluster.id).all()
    items = StationaryItem.query.order_by(
        StationaryItem.hsn_code.is_(None), StationaryItem.hsn_code, StationaryItem.name
    ).all()

    # ---- Sheet 1: Stationery Issued (per-issuance log) --------------------
    ws1 = wb.active
    ws1.title = 'Stationery Issued'
    ws1.cell(row=1, column=1, value=f'Stationery Issued — {period}').font = S['title_font']
    ws1.merge_cells('A1:E1')
    headers1 = ['Date', 'User', 'User Department', 'Item Issued', 'Qty']
    for i, label in enumerate(headers1, start=1):
        head(ws1.cell(row=2, column=i, value=label))
    r = 3
    issued_total = 0
    for req, line, item in rows:
        qty = line.quantity_fulfilled or 0
        issued_total += qty
        box(ws1.cell(row=r, column=1, value=req.fulfilled_at.strftime('%d/%m/%Y') if req.fulfilled_at else ''), align='center')
        box(ws1.cell(row=r, column=2, value=req.teacher.name if req.teacher else ''))
        box(ws1.cell(row=r, column=3, value=req.teacher.cluster.name if req.teacher and req.teacher.cluster else ''))
        box(ws1.cell(row=r, column=4, value=item.name))
        box(ws1.cell(row=r, column=5, value=qty), numfmt=_FMT_INT, align='center')
        r += 1
    box(ws1.cell(row=r, column=1, value='TOTAL'), bold=True, fill='total')
    for c in (2, 3, 4):
        box(ws1.cell(row=r, column=c), bold=True, fill='total')
    box(ws1.cell(row=r, column=5, value=issued_total), numfmt=_FMT_INT, bold=True, fill='total', align='center')
    for i, wdt in enumerate([13, 26, 24, 32, 9], start=1):
        ws1.column_dimensions[get_column_letter(i)].width = wdt
    ws1.freeze_panes = 'A3'
    ws1.auto_filter.ref = f'A2:E{r - 1}'

    # qty_id[item_id][cluster_id] within the filtered window (shared by sheets 2-4)
    dept_q = (
        db.session.query(RequestLine.item_id, Teacher.cluster_id, func.sum(RequestLine.quantity_fulfilled))
        .select_from(Request)
        .join(RequestLine, RequestLine.request_id == Request.id)
        .join(Teacher, Request.teacher_id == Teacher.id)
        .filter(Request.status == 'fulfilled')
    )
    dept_q = _apply_window(dept_q, filters)
    dept_q = dept_q.group_by(RequestLine.item_id, Teacher.cluster_id)
    qty_id = {}
    for item_id, cluster_id, q in dept_q.all():
        qty_id.setdefault(item_id, {})[cluster_id] = int(q or 0)

    # ---- Sheet 2: Departmental Expenditure (items x departments) ----------
    ws2 = wb.create_sheet('Departmental Expenditure')
    ws2.cell(row=1, column=1, value=f'Departmental Stationery Expenditure — {period}').font = S['title_font']
    fixed2 = ['HSN Code', 'ITEMS', 'Unit', 'Item Price (₹)', 'TOTAL QTY.']
    for i, label in enumerate(fixed2, start=1):
        ws2.cell(row=2, column=i, value=label)
        ws2.merge_cells(start_row=2, start_column=i, end_row=3, end_column=i)
        head(ws2.cell(row=2, column=i))
        head(ws2.cell(row=3, column=i))
    col = len(fixed2) + 1
    dept_cols = []  # (cluster_id, qty_col)
    for dept in departments:
        ws2.cell(row=2, column=col, value=dept.name)
        ws2.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)
        ws2.cell(row=3, column=col, value='QTY.')
        ws2.cell(row=3, column=col + 1, value='TOTAL (₹)')
        dept_cols.append((dept.id, col))
        col += 2
    gcol = col
    ws2.cell(row=2, column=gcol, value='G. Total')
    ws2.merge_cells(start_row=2, start_column=gcol, end_row=2, end_column=gcol + 1)
    ws2.cell(row=3, column=gcol, value='QTY.')
    ws2.cell(row=3, column=gcol + 1, value='TOTAL (₹)')
    ncols2 = gcol + 1
    for c in range(len(fixed2) + 1, ncols2 + 1):
        head(ws2.cell(row=2, column=c))
        head(ws2.cell(row=3, column=c))

    r = 4
    col_totals = {}  # column -> running total (qty cols + money cols + grand)
    for item in items:
        rate = rates.get(item.id, 0.0) or 0.0
        per_dept = qty_id.get(item.id, {})
        total_qty = sum(per_dept.values())
        box(ws2.cell(row=r, column=1, value=item.hsn_code or ''), align='center')
        box(ws2.cell(row=r, column=2, value=item.name))
        box(ws2.cell(row=r, column=3, value=item.unit), align='center')
        box(ws2.cell(row=r, column=4, value=round(rate, 2)), numfmt=_FMT_MONEY, fill='cream', align='center')
        box(ws2.cell(row=r, column=5, value=total_qty), numfmt=_FMT_INT, align='center')
        for cluster_id, qcol in dept_cols:
            q = per_dept.get(cluster_id, 0)
            amt = round(q * rate, 2)
            box(ws2.cell(row=r, column=qcol, value=q), numfmt=_FMT_INT, align='center')
            box(ws2.cell(row=r, column=qcol + 1, value=amt), numfmt=_FMT_MONEY, align='center')
            col_totals[qcol] = col_totals.get(qcol, 0) + q
            col_totals[qcol + 1] = col_totals.get(qcol + 1, 0) + amt
        gtotal = round(total_qty * rate, 2)
        box(ws2.cell(row=r, column=gcol, value=total_qty), numfmt=_FMT_INT, align='center')
        box(ws2.cell(row=r, column=gcol + 1, value=gtotal), numfmt=_FMT_MONEY, align='center')
        col_totals[5] = col_totals.get(5, 0) + total_qty
        col_totals[gcol] = col_totals.get(gcol, 0) + total_qty
        col_totals[gcol + 1] = col_totals.get(gcol + 1, 0) + gtotal
        r += 1
    # TOTAL row
    box(ws2.cell(row=r, column=1, value='TOTAL'), bold=True, fill='total')
    for c in (2, 3, 4):
        box(ws2.cell(row=r, column=c), bold=True, fill='total')
    for c in range(5, ncols2 + 1):
        fmt = _FMT_MONEY if (c == gcol + 1 or (c > 5 and (c - len(fixed2)) % 2 == 0)) else _FMT_INT
        box(ws2.cell(row=r, column=c, value=col_totals.get(c, 0)), numfmt=fmt, bold=True, fill='total', align='center')
    ws2.column_dimensions['A'].width = 11
    ws2.column_dimensions['B'].width = 32
    ws2.column_dimensions['C'].width = 7
    ws2.column_dimensions['D'].width = 13
    ws2.column_dimensions['E'].width = 11
    ws2.freeze_panes = 'F4'

    # ---- Sheet 3: Departmental Detail (one stacked block per department) ---
    ws4 = wb.create_sheet('Departmental Detail')
    ws4.cell(row=1, column=1, value=f'Departmental Detail — {period}').font = S['title_font']
    rr = 3
    det_headers = ['HSN Code', 'ITEMS', 'Unit', 'Item Price (₹)', 'QTY.', 'TOTAL (₹)']
    for dept in departments:
        # Items this department actually drew in the window
        dept_items = [(it, qty_id.get(it.id, {}).get(dept.id, 0)) for it in items]
        dept_items = [(it, q) for it, q in dept_items if q > 0]
        if not dept_items:
            continue
        ws4.cell(row=rr, column=1, value=dept.name)
        ws4.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=6)
        tc = ws4.cell(row=rr, column=1)
        tc.font = S['bold']
        tc.alignment = S['center']
        tc.fill = S['total']
        for c in range(1, 7):
            ws4.cell(row=rr, column=c).border = S['border']
        rr += 1
        for i, label in enumerate(det_headers, start=1):
            head(ws4.cell(row=rr, column=i, value=label))
        rr += 1
        block_total = 0.0
        for it, q in dept_items:
            rate = rates.get(it.id, 0.0) or 0.0
            amt = round(q * rate, 2)
            block_total += amt
            box(ws4.cell(row=rr, column=1, value=it.hsn_code or ''), align='center')
            box(ws4.cell(row=rr, column=2, value=it.name))
            box(ws4.cell(row=rr, column=3, value=it.unit), align='center')
            box(ws4.cell(row=rr, column=4, value=round(rate, 2)), numfmt=_FMT_MONEY, fill='cream', align='center')
            box(ws4.cell(row=rr, column=5, value=q), numfmt=_FMT_INT, align='center')
            box(ws4.cell(row=rr, column=6, value=amt), numfmt=_FMT_MONEY, align='center')
            rr += 1
        box(ws4.cell(row=rr, column=1, value='TOTAL'), bold=True, fill='total')
        for c in (2, 3, 4, 5):
            box(ws4.cell(row=rr, column=c), bold=True, fill='total')
        box(ws4.cell(row=rr, column=6, value=round(block_total, 2)), numfmt=_FMT_MONEY, bold=True, fill='total', align='center')
        rr += 2  # gap before next block
    for colL, wdt in zip('ABCDEF', [11, 34, 7, 13, 9, 13]):
        ws4.column_dimensions[colL].width = wdt

    # ---- Sheet 4: Department Summary (spend per department) ----------------
    ws5 = wb.create_sheet('Department Summary')
    ws5.cell(row=1, column=1, value=f'Stationery Details — {period}').font = S['title_font']
    ws5.merge_cells('A1:B1')
    for i, label in enumerate(['Department', 'Total Expenditure (₹)'], start=1):
        head(ws5.cell(row=2, column=i, value=label))
    dept_spend = []
    for dept in departments:
        spend = 0.0
        for it in items:
            q = qty_id.get(it.id, {}).get(dept.id, 0)
            if q:
                spend += q * (rates.get(it.id, 0.0) or 0.0)
        dept_spend.append((dept.name, round(spend, 2)))
    dept_spend.sort(key=lambda x: x[1], reverse=True)
    rr = 3
    grand = 0.0
    for name, spend in dept_spend:
        grand += spend
        box(ws5.cell(row=rr, column=1, value=name))
        box(ws5.cell(row=rr, column=2, value=spend), numfmt=_FMT_MONEY, align='center')
        rr += 1
    box(ws5.cell(row=rr, column=1, value='TOTAL'), bold=True, fill='total')
    box(ws5.cell(row=rr, column=2, value=round(grand, 2)), numfmt=_FMT_MONEY, bold=True, fill='total', align='center')
    ws5.column_dimensions['A'].width = 34
    ws5.column_dimensions['B'].width = 22
    ws5.freeze_panes = 'A3'

    # ---- Sheet 5: MIS (items x months, Qty + Amount per month) ------------
    year = _report_year(filters)
    mis_q = (
        db.session.query(RequestLine.item_id, func.extract('month', Request.fulfilled_at), func.sum(RequestLine.quantity_fulfilled))
        .select_from(Request)
        .join(RequestLine, RequestLine.request_id == Request.id)
        .filter(Request.status == 'fulfilled')
        .filter(func.extract('year', Request.fulfilled_at) == year)
        .group_by(RequestLine.item_id, func.extract('month', Request.fulfilled_at))
    )
    qty_month = {}
    for item_id, m, q in mis_q.all():
        qty_month.setdefault(item_id, {})[int(m)] = int(q or 0)

    ws3 = wb.create_sheet('MIS')
    ws3.cell(row=1, column=1, value=f'STATIONERY MIS {year}').font = S['title_font']
    fixed3 = ['HSN Code', 'ITEMS', 'Unit', 'Item Price (₹)']
    for i, label in enumerate(fixed3, start=1):
        ws3.cell(row=2, column=i, value=label)
        ws3.merge_cells(start_row=2, start_column=i, end_row=3, end_column=i)
        head(ws3.cell(row=2, column=i))
        head(ws3.cell(row=3, column=i))
    col = len(fixed3) + 1
    month_cols = []  # (month_index, qty_col)
    for mi, mn in enumerate(MONTHS_SHORT, start=1):
        ws3.cell(row=2, column=col, value=mn.upper())
        ws3.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)
        ws3.cell(row=3, column=col, value='Qty.')
        ws3.cell(row=3, column=col + 1, value='Amount (₹)')
        month_cols.append((mi, col))
        col += 2
    ncols3 = col - 1
    for c in range(len(fixed3) + 1, ncols3 + 1):
        head(ws3.cell(row=2, column=c))
        head(ws3.cell(row=3, column=c))

    r = 4
    mis_totals = {}
    for item in items:
        rate = rates.get(item.id, 0.0) or 0.0
        per_month = qty_month.get(item.id, {})
        box(ws3.cell(row=r, column=1, value=item.hsn_code or ''), align='center')
        box(ws3.cell(row=r, column=2, value=item.name))
        box(ws3.cell(row=r, column=3, value=item.unit), align='center')
        box(ws3.cell(row=r, column=4, value=round(rate, 2)), numfmt=_FMT_MONEY, fill='cream', align='center')
        for mi, qcol in month_cols:
            q = per_month.get(mi, 0)
            amt = round(q * rate, 2)
            box(ws3.cell(row=r, column=qcol, value=q), numfmt=_FMT_INT, align='center')
            box(ws3.cell(row=r, column=qcol + 1, value=amt), numfmt=_FMT_MONEY, align='center')
            mis_totals[qcol] = mis_totals.get(qcol, 0) + q
            mis_totals[qcol + 1] = mis_totals.get(qcol + 1, 0) + amt
        r += 1
    box(ws3.cell(row=r, column=1, value='TOTAL'), bold=True, fill='total')
    for c in (2, 3, 4):
        box(ws3.cell(row=r, column=c), bold=True, fill='total')
    for mi, qcol in month_cols:
        box(ws3.cell(row=r, column=qcol, value=mis_totals.get(qcol, 0)), numfmt=_FMT_INT, bold=True, fill='total', align='center')
        box(ws3.cell(row=r, column=qcol + 1, value=mis_totals.get(qcol + 1, 0)), numfmt=_FMT_MONEY, bold=True, fill='total', align='center')
    ws3.column_dimensions['A'].width = 11
    ws3.column_dimensions['B'].width = 32
    ws3.column_dimensions['C'].width = 7
    ws3.column_dimensions['D'].width = 13
    ws3.freeze_panes = 'E4'

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
