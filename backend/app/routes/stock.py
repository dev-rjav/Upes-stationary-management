from flask import Blueprint, jsonify, request
from flask_login import current_user, login_required

from app import db
from app.models import Item, Stock, StockIn
from app.services import stock_service

bp = Blueprint("stock", __name__, url_prefix="/api/stock")


@bp.get("")
@login_required
def list_stock():
    q = request.args.get("q", "").strip().lower()
    low_only = request.args.get("low") == "1"
    rows = (
        db.session.query(Stock, Item)
        .join(Item, Item.id == Stock.item_id)
        .order_by(Item.name)
        .all()
    )
    out = []
    for s, i in rows:
        if not i.active:
            continue
        d = s.to_dict()
        if q and q not in i.name.lower():
            continue
        if low_only and not d["is_low"]:
            continue
        out.append(d)
    return jsonify(out)


@bp.post("/in")
@login_required
def stock_in():
    d = request.get_json(silent=True) or {}
    item_id = d.get("item_id")
    qty = int(d.get("qty") or 0)
    if not item_id or qty <= 0:
        return jsonify({"error": "item_id and positive qty required"}), 400
    stock_service.add_stock(item_id, qty, float(d.get("unit_rate") or 0), source=d.get("source", "manual"), note=d.get("note", ""), created_by=current_user.name)
    db.session.commit()
    s = Stock.query.get(item_id)
    return jsonify(s.to_dict()), 201


@bp.post("/import")
@login_required
def stock_import():
    """D14: vendor Excel stock-in — fuzzy match + alias learning."""
    rows = request.get_json(silent=True) or {}
    parsed = rows.get("rows") or []
    import openpyxl
    from thefuzz import fuzz
    from werkzeug.utils import secure_filename

    matched, unmatched = [], []
    for r in parsed:
        name = str(r.get("item") or "").strip()
        qty = r.get("qty")
        rate = r.get("rate")
        if not name or qty in (None, ""):
            continue
        try:
            qty = int(float(str(qty)))
        except ValueError:
            continue
        item = Item.find_by_alias(name) or Item.query.filter(db.func.lower(Item.name) == name.lower()).first()
        method = "alias" if item else None
        if not item:
            best, best_score = None, 0
            for it in Item.query.filter_by(active=True).all():
                sc = max(fuzz.ratio(name.lower(), it.name.lower()), fuzz.partial_ratio(name.lower(), it.name.lower()) * 0.9)
                if sc > best_score:
                    best, best_score = it, sc
            if best_score >= 75:
                item = best
                method = "fuzzy"
                if r.get("learn", True):
                    from app.models import ItemAlias
                    if not Item.find_by_alias(name):
                        db.session.add(ItemAlias(alias=name, item_id=item.id, source="learned"))
                        db.session.flush()
        if item:
            matched.append({"alias": name, "item_id": item.id, "item": item.name, "qty": qty, "rate": float(rate or item.rate or 0), "method": method, "score": best_score if method == "fuzzy" else None})
        else:
            unmatched.append({"alias": name, "qty": qty, "rate": rate})
    if rows.get("apply", False):
        for m in matched:
            stock_service.add_stock(m["item_id"], m["qty"], m["rate"], source="excel", created_by=current_user.name)
        db.session.commit()
    return jsonify({"matched": matched, "unmatched": unmatched, "applied": bool(rows.get("apply"))})


@bp.get("/history")
@login_required
def history():
    per = int(request.args.get("per_page", 30))
    rows = StockIn.query.order_by(StockIn.created_at.desc()).limit(per).all()
    return jsonify([r.to_dict() for r in rows])


@bp.post("/import/parse")
@login_required
def import_parse():
    """Parse an uploaded vendor Excel: expects Item | Qty | Rate columns (any header row in top 10)."""
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "file required"}), 400
    import io
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
    ws = wb.active
    grid = [list(r) for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 2000), values_only=True)]
    header_i = None
    for i in range(10):
        if i < len(grid) and any(grid[i][j] and str(grid[i][j]).strip().lower() in ("item", "item name", "items", "item issued") for j in range(min(6, len(grid[i])))):
            header_i = i
            break
    rows = []
    if header_i is not None:
        header = [str(c).strip().lower() if c else "" for c in grid[header_i]]
        def col(*keys):
            for k in keys:
                if k in header:
                    return header.index(k)
            return None
        ic, qc, rc = col("item", "item name", "items"), col("qty", "quantity"), col("rate", "item price", "price")
        if ic is None:
            ic = 0
        for r in grid[header_i + 1:]:
            if len(r) <= max(x or 0 for x in (ic, qc, rc) if x is not None):
                continue
            name = r[ic]
            if not name:
                continue
            qty = r[qc] if qc is not None else None
            rate = r[rc] if rc is not None else None
            if qty in (None, ""):
                continue
            rows.append({"item": str(name).strip(), "qty": qty, "rate": rate})
    wb.close()
    return jsonify({"rows": rows})